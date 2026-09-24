"""Cria meridian_finance e carrega empresa, conta, plano e fornecedor.

Lê o .env do Meridian só para achar a instância. Nunca grava em vision_check.
Os arquivos-fonte ficam em Dados Grupo Alvim/ e não entram no Git.
"""

from __future__ import annotations

import os
import re
import unicodedata
from pathlib import Path

import openpyxl
import psycopg
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "Dados Grupo Alvim" / "00. DADOS LOJAS- GRUPO ALVIM"
SQL = ROOT / "db" / "001_nucleo.sql"
MERIDIAN_ENV = Path(__file__).resolve().parents[2] / "Check_visaodono" / "backend" / ".env"


def load_env(path: Path) -> dict[str, str]:
    env: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        env[key.strip()] = value.strip()
    return env


def norm(value) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.upper().split())


def digits(value) -> str:
    return re.sub(r"\D", "", "" if value is None else str(value))


def split_conta(value) -> tuple[str, str | None]:
    text = "" if value is None else str(value).strip()
    if "-" in text:
        numero, digito = text.rsplit("-", 1)
        return numero.lstrip("0") or "0", digito.strip() or None
    return text.lstrip("0") or "0", None


def apelido_de(razao: str) -> str:
    nome = norm(razao)
    if "SUPER KING" in nome:
        return "SUPER KING"
    if nome.startswith("KING "):
        return "KING"
    if "PARTICIPA" in nome:
        return "ALVIM PARTICIPACOES"
    return nome.split()[0]


def sheet_rows(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def find_file(prefix: str) -> Path:
    for path in DATA.iterdir():
        if path.name.upper().startswith(prefix) and path.suffix.lower() == ".xlsx":
            return path
    raise FileNotFoundError(prefix)


def main() -> None:
    env = load_env(MERIDIAN_ENV)
    host = env["DB_HOST"]
    user = env["DB_USER"]
    password = env["DB_PASS"]
    port = int(env.get("DB_PORT") or 5432)
    base = dict(host=host, user=user, password=password, port=port, connect_timeout=15)

    with psycopg.connect(**base, dbname="postgres", autocommit=True) as conn:
        exists = conn.execute("select 1 from pg_database where datname = 'meridian_finance'").fetchone()
        if not exists:
            conn.execute("create database meridian_finance")
            print("database criado")
        else:
            print("database ja existe")

    with psycopg.connect(**base, dbname="meridian_finance", autocommit=True) as conn:
        pronto = conn.execute("select to_regclass('public.empresas')").fetchone()[0]
        if not pronto:
            conn.execute(SQL.read_text(encoding="utf-8"))
            print("schema aplicado")
        else:
            print("schema ja existe")

    lojas = []
    for row in sheet_rows(find_file("DADOS LOJAS ATUALIZADO- 2026"))[4:]:
        if not row or not row[1]:
            continue
        cnpj = digits(row[3])
        bk = digits(row[2]) or None
        lojas.append({
            "razao": str(row[1]).strip(),
            "apelido": apelido_de(row[1]),
            "bk": bk,
            "cnpj": cnpj or None,
            "ie": (str(row[4]).strip() if row[4] else None),
            "endereco": (str(row[5]).strip() if row[5] else None),
            "cidade": (str(row[6]).strip() if row[6] else None),
            "cep": (str(row[7]).strip() if row[7] else None),
            "tipo": "loja" if bk else "holding",
        })

    pdf = next(DATA.glob("CONTAS BANCO*.pdf"))
    texto = "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf)).pages)
    bb = []
    for linha in texto.splitlines():
        partes = linha.split()
        if len(partes) < 4 or not digits(partes[-1]):
            continue
        cnpj = digits(partes[-1])
        if len(cnpj) != 14:
            continue
        conta, agencia = partes[-2], partes[-3]
        apelido = " ".join(partes[:-3]).strip()
        if not apelido or apelido.upper().startswith("LOJA"):
            continue
        numero, digito = split_conta(conta)
        bb.append({"apelido": apelido, "cnpj": cnpj, "agencia": agencia, "numero": numero, "digito": digito})

    itau = []
    for row in sheet_rows(find_file("DADOS ITA"))[4:]:
        if not row or not row[1] or not row[4]:
            continue
        numero, digito = split_conta(row[4])
        itau.append({
            "razao": str(row[1]).strip(),
            "cnpj": digits(row[2]) or None,
            "agencia": str(row[3]).strip(),
            "numero": numero,
            "digito": digito,
        })

    wb = openpyxl.load_workbook(find_file("PLANODECONTAS"), read_only=True, data_only=True)
    plano_ws = wb["Plano de Contas"]
    planos = []
    for row in list(plano_ws.iter_rows(values_only=True))[2:]:
        if not row or not row[0]:
            continue
        tipo = norm(row[1])
        natureza = norm(row[4]) if len(row) > 4 and row[4] else ""
        planos.append({
            "nome": str(row[0]).strip(),
            "tipo": "a_receber" if "RECEBER" in tipo else "a_pagar",
            "natureza": "fixa" if natureza == "FIXA" else "variavel" if natureza == "VARIAVEL" else None,
        })
    wb.close()

    fornec = []
    for row in sheet_rows(DATA / "Dados fornecedor.xlsx")[2:]:
        if not row or not row[3]:
            continue
        doc = digits(row[1]) or None
        tipo_pessoa = "cpf" if doc and len(doc) <= 11 else "cnpj" if doc else None
        fornec.append({
            "nome": str(row[3]).strip(),
            "razao": (str(row[6]).strip() if row[6] else None),
            "doc": doc,
            "tipo_pessoa": tipo_pessoa,
            "plano": (str(row[15]).strip() if row[15] else None),
            "logradouro": (str(row[8]).strip() if row[8] else None),
            "numero": (str(row[9]).strip() if row[9] else None),
            "bairro": (str(row[11]).strip() if row[11] else None),
            "cidade": (str(row[12]).strip() if row[12] else None),
            "estado": (str(row[13]).strip() if row[13] else None),
            "cep": (str(row[14]).strip() if row[14] else None),
            "pix": (str(row[22]).strip() if row[22] else None),
            "tipo_pix": (str(row[23]).strip() if row[23] else None),
            "cod_banco": (str(row[24]).strip() if row[24] else None),
            "banco": (str(row[25]).strip() if row[25] else None),
            "agencia": (str(row[26]).strip() if row[26] else None),
            "digito_agencia": (str(row[27]).strip() if row[27] else None),
            "conta": (str(row[28]).strip() if row[28] else None),
            "digito_conta": (str(row[29]).strip() if len(row) > 29 and row[29] else None),
        })

    with psycopg.connect(**base, dbname="meridian_finance") as conn:
        with conn.cursor() as cur:
            cur.execute("""
                truncate fornecedor_pagamentos, despesas, fornecedores,
                         plano_contas, contas_bancarias, empresas restart identity cascade
            """)
            for loja in lojas:
                cur.execute(
                    """
                    insert into empresas (apelido, razao_social, bk_number, cnpj, inscricao_estadual, endereco, cidade, cep, tipo)
                    values (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (loja["apelido"], loja["razao"], loja["bk"], loja["cnpj"], loja["ie"], loja["endereco"], loja["cidade"], loja["cep"], loja["tipo"]),
                )
            cur.execute("select id, cnpj, apelido from empresas")
            por_cnpj = {cnpj: (eid, apelido) for eid, cnpj, apelido in cur.fetchall() if cnpj}

            for conta in bb:
                empresa = por_cnpj.get(conta["cnpj"])
                if not empresa:
                    print("bb sem empresa", conta["apelido"])
                    continue
                cur.execute(
                    """
                    update empresas set apelido = %s where id = %s and apelido is distinct from %s
                    """,
                    (conta["apelido"], empresa[0], conta["apelido"]),
                )
                cur.execute(
                    """
                    insert into contas_bancarias (empresa_id, banco, agencia, numero, digito)
                    values (%s, 'banco_do_brasil', %s, %s, %s)
                    on conflict (banco, agencia, numero) do nothing
                    """,
                    (empresa[0], conta["agencia"], conta["numero"], conta["digito"]),
                )

            for conta in itau:
                empresa = por_cnpj.get(conta["cnpj"]) if conta["cnpj"] else None
                if not empresa:
                    print("itau sem empresa", conta["razao"])
                    continue
                cur.execute(
                    """
                    insert into contas_bancarias (empresa_id, banco, agencia, numero, digito)
                    values (%s, 'itau', %s, %s, %s)
                    on conflict (banco, agencia, numero) do nothing
                    """,
                    (empresa[0], conta["agencia"], conta["numero"], conta["digito"]),
                )

            for plano in planos:
                cur.execute(
                    """
                    insert into plano_contas (nome, tipo, natureza)
                    values (%s, %s, %s)
                    on conflict (nome) do nothing
                    """,
                    (plano["nome"], plano["tipo"], plano["natureza"]),
                )
            cur.execute("select id, nome from plano_contas")
            por_plano = {norm(nome): eid for eid, nome in cur.fetchall()}

            sem_plano = 0
            vistos = set()
            for item in fornec:
                if item["doc"] and item["doc"] in vistos:
                    continue
                if item["doc"]:
                    vistos.add(item["doc"])
                plano_id = por_plano.get(norm(item["plano"])) if item["plano"] else None
                if item["plano"] and not plano_id:
                    sem_plano += 1
                cur.execute(
                    """
                    insert into fornecedores (
                      nome, razao_social, cpf_cnpj, tipo_pessoa, plano_conta_id,
                      logradouro, numero, bairro, cidade, estado, cep
                    ) values (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    returning id
                    """,
                    (item["nome"], item["razao"], item["doc"], item["tipo_pessoa"], plano_id,
                     item["logradouro"], item["numero"], item["bairro"], item["cidade"], item["estado"], item["cep"]),
                )
                fid = cur.fetchone()[0]
                if item["pix"]:
                    cur.execute(
                        """
                        insert into fornecedor_pagamentos (fornecedor_id, meio, chave_pix, tipo_chave_pix)
                        values (%s, 'pix', %s, %s)
                        """,
                        (fid, item["pix"], item["tipo_pix"]),
                    )
                if item["conta"]:
                    cur.execute(
                        """
                        insert into fornecedor_pagamentos (
                          fornecedor_id, meio, codigo_banco, banco, agencia, digito_agencia, numero, digito
                        ) values (%s, 'conta', %s, %s, %s, %s, %s, %s)
                        """,
                        (fid, item["cod_banco"], item["banco"], item["agencia"], item["digito_agencia"], item["conta"], item["digito_conta"]),
                    )

        conn.commit()
        with conn.cursor() as cur:
            for tabela in ("empresas", "contas_bancarias", "plano_contas", "fornecedores", "fornecedor_pagamentos"):
                cur.execute(f"select count(*) from {tabela}")
                print(tabela, cur.fetchone()[0])
            cur.execute("select banco, count(*) from contas_bancarias group by banco order by banco")
            print("contas", cur.fetchall())
            cur.execute("select tipo, count(*) from empresas group by tipo order by tipo")
            print("empresas_tipo", cur.fetchall())
            print("fornecedor_sem_plano", sem_plano)


if __name__ == "__main__":
    main()
